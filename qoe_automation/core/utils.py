import pandas as pd
import pdfplumber
from datetime import datetime, timedelta
from decimal import Decimal
import re
import io
import os
import openpyxl
from .models import QuickbookRecord, PayrollRecord, Adjustment, StocAccountingData, ProfitLossRecord, DataSource, SupplementalNote, RecastPL
from openpyxl import Workbook
import openpyxl.styles

def process_quickbooks_file(data_source):
    file_path = data_source.file_path.path
    processed_count = 0
    try:
        file_extension = os.path.splitext(file_path)[1].lower()
        df = None

        if file_extension == '.csv':
            df = pd.read_csv(file_path, encoding='latin-1', sep=None, engine='python', header=None, on_bad_lines='warn')
            header_row_index = -1
            max_found_headers = 0
            expected_gl_headers = [
                "Distribution account", "Transaction date", "Transaction type",
                "Num", "Name", "Memo/Description", "Split account", "Amount", "Balance"
            ]

            for i in range(min(df.shape[0], 15)):
                current_row_values = [str(x).strip() for x in df.iloc[i].dropna().tolist()]
                row_content_lower = [x.lower() for x in current_row_values]

                found_headers_count = 0
                for expected_h in expected_gl_headers:
                    if expected_h.lower() in row_content_lower:
                        found_headers_count += 1

                if found_headers_count > max_found_headers:
                    max_found_headers = found_headers_count
                    header_row_index = i

                if max_found_headers >= len(expected_gl_headers) - 2:
                    break

            if header_row_index == -1:
                raise ValueError(f"QuickBooks GL header not found in the CSV file. Please ensure '{', '.join(expected_gl_headers)}' columns are present in at least one row within the first 15 rows.")

            df.columns = [str(col).strip() if pd.notna(col) else f"Unnamed_{j}" for j, col in enumerate(df.iloc[header_row_index])]
            df = df[header_row_index + 1:].reset_index(drop=True)

        elif file_extension in ['.xlsx', '.xls']:
            df = pd.read_excel(file_path, header=4) 
        else:
            raise ValueError(f"Unsupported file type for QuickBooks: {file_extension}. Only .csv and .xlsx/.xls are supported.")

        if df is None or df.empty:
            raise ValueError("Could not read QuickBooks file or file is empty.")

        column_mapping = {
            'distribution account': 'Distribution account',
            'transaction date': 'Transaction date',
            'transaction type': 'Transaction type',
            'num': 'Num',
            'name': 'Name',
            'memo/description': 'Memo/Description',
            'split account': 'Split account',
            'amount': 'Amount',
            'balance': 'Balance'
        }
        
        cleaned_df_columns = {col: str(col).strip().lower() for col in df.columns}
        new_columns = {}
        for old_col_name, cleaned_old_col_name in cleaned_df_columns.items():
            if cleaned_old_col_name in column_mapping:
                new_columns[old_col_name] = column_mapping[cleaned_old_col_name]
            else:
                new_columns[old_col_name] = old_col_name

        df.rename(columns=new_columns, inplace=True)
        
        required_cols_for_processing = ['Transaction date', 'Distribution account', 'Amount', 'Balance']
        if not all(col in df.columns for col in required_cols_for_processing):
            missing_cols = [col for col in required_cols_for_processing if col not in df.columns]
            raise ValueError(f"Missing critical columns after processing QuickBooks file: {', '.join(missing_cols)}. Available columns: {df.columns.tolist()}")

        QuickbookRecord.objects.filter(data_source=data_source).delete()

        for _, row in df.iterrows():
            if pd.isna(row['Transaction date']) or \
               'Beginning Balance' in str(row['Distribution account']) or \
               'Total for' in str(row['Distribution account']) or \
               pd.isna(row['Amount']):
                continue

            amount_str = str(row['Amount']).strip()
            amount_str = amount_str.replace('$', '').replace(',', '')
            if amount_str.startswith('(') and amount_str.endswith(')'):
                amount_str = '-' + amount_str[1:-1]
            
            if not amount_str:
                amount_str = '0'

            amount = Decimal(amount_str)
            debit = Decimal('0')
            credit = Decimal('0')

            if amount < 0:
                debit = abs(amount)
            else:
                credit = amount

            try:
                record_date = pd.to_datetime(row['Transaction date']).date()
            except ValueError:
                continue

            account_full_name = str(row['Distribution account'])
            account_number_match = re.match(r'(\d+)\s(.*)', account_full_name)

            account_number = ''
            account_name = account_full_name

            if account_number_match:
                account_number = account_number_match.group(1)
                account_name = account_number_match.group(2)
            else:
                account_number = ''
                account_name = account_full_name

            QuickbookRecord.objects.create(
                data_source=data_source,
                date=record_date,
                account_number=account_number,
                account_name=account_name,
                description=row['Memo/Description'] if pd.notna(row['Memo/Description']) else (row['Name'] if pd.notna(row['Name']) else row['Transaction type']),
                debit=debit,
                credit=credit,
                balance=Decimal(str(row['Balance']).strip().replace('$', '').replace(',', '').replace('(', '-').replace(')', '') if pd.notna(row['Balance']) else '0')
            )
            processed_count += 1
        
        data_source.processed_records_count = processed_count
        data_source.save()

    except Exception as e:
        raise Exception(f"Error processing QuickBooks file for data source {data_source.id}: {str(e)}")

def process_payroll_file(data_source):
    file_path = data_source.file_path.path
    processed_count = 0
    try:
        file_extension = os.path.splitext(file_path)[1].lower()
        df = None

        if file_extension == '.csv':
            df = pd.read_csv(file_path, encoding='latin-1', on_bad_lines='warn', header=None)
        elif file_extension in ['.xlsx', '.xls']:
            df = pd.read_excel(file_path, header=None)
        else:
            raise ValueError(f"Unsupported file type for payroll: {file_extension}. Only .csv and .xlsx/.xls are supported.")

        if df is None or df.empty:
            raise ValueError("Could not read payroll file or file is empty.")

        header_row_index = -1
        expected_headers = ["Employee Name", "TIN", "Pay Frequency", "Department"]

        for i, row in df.iterrows():
            current_row_values = [str(x).strip() for x in row.dropna().tolist()]
            row_content_str = " ".join(current_row_values).lower()
            if all(h.lower() in row_content_str for h in expected_headers):
                header_row_index = i
                break

        if header_row_index == -1:
            raise ValueError(f"Payroll header not found in the file. Please ensure '{', '.join(expected_headers)}' columns are present.")

        df.columns = [str(col).strip() if pd.notna(col) else f"Unnamed_{j}" for j, col in enumerate(df.iloc[header_row_index])]
        df = df[header_row_index + 1:].reset_index(drop=True)

        seen_columns = {}
        unique_cleaned_header_columns = []
        for col in df.columns:
            original_col = col
            count = seen_columns.get(original_col, 0)
            if count > 0:
                col = f'{original_col}.{count}'
            while col in unique_cleaned_header_columns:
                count += 1
                col = f'{original_col}.{count}'
            unique_cleaned_header_columns.append(col)
            seen_columns[original_col] = count + 1
        
        df.columns = unique_cleaned_header_columns

        PayrollRecord.objects.filter(data_source=data_source).delete()

        if 'Employee Name' not in df.columns:
            raise ValueError("Processed DataFrame does not contain 'Employee Name' column after header detection.")
        df['Employee Name'] = df['Employee Name'].astype(str)

        df_filtered = df.loc[df['Employee Name'].notna() & ~df['Employee Name'].str.contains('Total', case=False, na=False)].copy()

        earning_amount_cols = ['Amount'] + [f'Amount.{i}' for i in range(1, 19)]
        earning_amount_cols = [col for col in earning_amount_cols if col in df_filtered.columns]

        benefit_cols = [
            '401k Employee Contribution',
            'Pre-tax Medical Health Insurance Deduc',
            'AFLAC pre-tax',
        ]
        benefit_cols = [col for col in benefit_cols if col in df_filtered.columns]
        
        for _, row in df_filtered.iterrows():
            employee_name = row.get('Employee Name')
            if pd.isna(employee_name):
                continue

            record_date = extract_date_from_payment_columns(row)
            if record_date is None:
                continue

            compensation = Decimal('0')
            for col in earning_amount_cols:
                if pd.notna(row[col]): 
                    try:
                        cleaned_val = str(row[col]).replace('$', '').replace(',', '').strip()
                        if cleaned_val:
                            compensation += Decimal(cleaned_val)
                    except (ValueError, TypeError, AttributeError) as e:
                        pass
            taxes_str = str(row['Total Taxes']).replace('$', '').replace(',', '').strip() if pd.notna(row['Total Taxes']) else '0'
            taxes = Decimal(taxes_str) if taxes_str else Decimal('0')

            benefits = Decimal('0')
            for col in benefit_cols:
                if pd.notna(row[col]): 
                    try:
                        cleaned_val = str(row[col]).replace('$', '').replace(',', '').strip()
                        if cleaned_val:
                            benefits += Decimal(cleaned_val)
                    except (ValueError, TypeError, AttributeError) as e:
                        pass 

            PayrollRecord.objects.create(
                data_source=data_source,
                employee_name=employee_name,
                title="Employee",
                compensation=compensation,
                taxes=taxes,
                benefits=benefits,
                date=record_date
            )
            processed_count += 1
        
        data_source.processed_records_count = processed_count
        data_source.save()

    except Exception as e:
        raise Exception(f"Error processing Payroll file for data source {data_source.id}: {str(e)}")

def process_profit_and_loss_file(data_source):
    file_path = data_source.file_path.path
    processed_count = 0
    try:
        file_extension = os.path.splitext(file_path)[1].lower()
        df = None

        if file_extension == '.csv':
            df = pd.read_csv(file_path, encoding='latin-1', on_bad_lines='warn', header=4)
        elif file_extension in ['.xlsx', '.xls']:
            df = pd.read_excel(file_path, header=4)
        else:
            raise ValueError(f"Unsupported file type for P&L: {file_extension}. Only .csv and .xlsx/.xls are supported.")

        if df is None or df.empty:
            raise ValueError("Could not read P&L file or file is empty.")

        df.rename(columns={df.columns[0]: 'Account'}, inplace=True)

        month_columns = []
        for col in df.columns:
            stripped_col = str(col).strip()
            try:
                datetime.strptime(stripped_col, '%B %Y')
                month_columns.append(col)
            except ValueError:
                pass
        
        if not month_columns:
            raise ValueError("No monthly P&L columns found (e.g., 'January 2023'). Please ensure your file is in the expected pivoted format.")

        pl_line_items_map = {
            'Revenue': 'Total for INCOME',
            'COGS': 'Total for DIRECT COSTS',
            'Gross Profit': 'Gross Profit',
            'Operating Expenses': 'Total for Expenses',
            'EBITDA': 'Net Operating Income',
            'Net Income': 'Net Income'
        }

        pl_data_rows = {}
        for pl_field, account_name in pl_line_items_map.items():
            row_data = df[df['Account'] == account_name]
            if not row_data.empty:
                pl_data_rows[pl_field] = row_data.iloc[0]
            else:
                if pl_field in ['Revenue', 'COGS', 'Gross Profit', 'Operating Expenses', 'EBITDA', 'Net Income']:
                     raise ValueError(f"Missing required P&L line item row: '{account_name}'")


        def clean_and_convert_decimal(value):
            if pd.isna(value):
                return Decimal('0.00')
            s_value = str(value).strip().replace('$', '').replace(',', '')
            if s_value.startswith('(') and s_value.endswith(')'):
                s_value = '-' + s_value[1:-1]
            if not s_value:
                return Decimal('0.00')
            try:
                return Decimal(s_value)
            except Exception:
                return Decimal('0.00') 

        ProfitLossRecord.objects.filter(data_source=data_source).delete()

        for month_col in month_columns:
            try:

                date_obj = datetime.strptime(month_col, '%B %Y').date()

                last_day = (date_obj.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
                record_date = last_day

                revenue = clean_and_convert_decimal(pl_data_rows['Revenue'].get(month_col))
                cogs = clean_and_convert_decimal(pl_data_rows['COGS'].get(month_col))
                gross_profit = clean_and_convert_decimal(pl_data_rows['Gross Profit'].get(month_col))
                operating_expenses = clean_and_convert_decimal(pl_data_rows['Operating Expenses'].get(month_col))
                ebitda = clean_and_convert_decimal(pl_data_rows['EBITDA'].get(month_col))
                net_income = clean_and_convert_decimal(pl_data_rows['Net Income'].get(month_col))
                
                ProfitLossRecord.objects.create(
                    data_source=data_source,
                    date=record_date,
                    revenue=revenue,
                    cogs=cogs,
                    gross_profit=gross_profit,
                    operating_expenses=operating_expenses,
                    ebitda=ebitda,
                    net_income=net_income
                )
                processed_count += 1
            except Exception as e:
                print(f"Skipping monthly record for '{month_col}' due to data conversion or parsing error: {e}")
                continue

        data_source.processed_records_count = processed_count
        data_source.save()

    except Exception as e:
        raise Exception(f"Error processing Profit and Loss file for data source {data_source.id}: {str(e)}")

def process_supplemental_notes(data_source):
    file_path = data_source.file_path.path
    processed_count = 0
    try:

        SupplementalNote.objects.filter(data_source=data_source).delete()

        SupplementalNote.objects.create(
            data_source=data_source,
            note_type="Owner Compensation",
            description="Simulated owner compensation identified from supplemental notes.",
            impact_amount=Decimal('50000.00')
        )
        processed_count += 1

        data_source.processed_records_count = processed_count
        data_source.save()

    except Exception as e:
        raise Exception(f"Error processing supplemental notes for data source {data_source.id}: {str(e)}")

def suggest_adjustments(data_source):

    gl_records = QuickbookRecord.objects.filter(data_source=data_source)

    payroll_records = PayrollRecord.objects.filter(data_source=data_source)

    pl_records = ProfitLossRecord.objects.filter(data_source=data_source)

    adjustments_suggested_count = 0
    suggested_adjustments_list = [] 
    for record in gl_records:
        if record.debit > Decimal('5000.00'):
            ai_reason = generate_ai_explanation("large debit", record.description)
            adjustment, created = Adjustment.objects.get_or_create(
            description=f"Large debit: {record.description}",
            adjustment_type='OWNER_COMP', 
            amount=record.debit,
            date=record.date,
            status='AI_SUGGESTED',
            rationale="Identified as a large debit transaction.",
            source_reference=f"QB GL Account: {record.account_name} ({record.account_number})",
            is_ai_suggested=True,
            ai_confidence_score=Decimal('0.75'),
            ai_suggestion_reason=ai_reason,
            created_by=data_source.user 
        )
            if created: 
                adjustments_suggested_count += 1
                suggested_adjustments_list.append(adjustment)


    for record in payroll_records:
        if record.compensation > Decimal('10000.00'):
            ai_reason = generate_ai_explanation("large payroll compensation", record.title)
            adjustment, created = Adjustment.objects.get_or_create(
            description=f"Large payroll compensation: {record.employee_name} - {record.compensation}",
            adjustment_type='DISCRETIONARY',  
            amount=record.compensation,
            date=record.date,
            status='AI_SUGGESTED',
            rationale="Identified as a large payroll compensation.",
            source_reference=f"Payroll Employee: {record.employee_name}",
            is_ai_suggested=True,
            ai_confidence_score=Decimal('0.80'),
            ai_suggestion_reason=ai_reason,
            created_by=data_source.user 
        )
            if created: 
                adjustments_suggested_count += 1
                suggested_adjustments_list.append(adjustment)

    
    for record in pl_records:
        if record.net_income < 0:
            ai_reason = generate_ai_explanation("negative net income", f"Net income: {record.net_income}")
            adjustment, created = Adjustment.objects.get_or_create(
            description=f"Negative net income: {record.net_income}",
            adjustment_type='NORMALIZATION',
            amount=abs(record.net_income),
            date=record.date,
            status='AI_SUGGESTED',
            rationale="Identified as a negative net income.",
            source_reference=f"PL Record on {record.date}",
            is_ai_suggested=True,
            ai_confidence_score=Decimal('0.80'),
            ai_suggestion_reason=ai_reason,
            created_by=data_source.user 
        )
            if created:
                adjustments_suggested_count += 1
                suggested_adjustments_list.append(adjustment)

        if record.operating_expenses > Decimal('200000.00'):
            ai_reason = generate_ai_explanation("high operating expenses", f"Operating expenses: {record.operating_expenses}")
            adjustment, created = Adjustment.objects.get_or_create(
            description=f"High operating expenses: {record.operating_expenses}",
            adjustment_type='NORMALIZATION',
            amount=record.operating_expenses,
            date=record.date,
            status='AI_SUGGESTED',
            rationale="Identified as high operating expenses.",
            source_reference=f"PL Record on {record.date}",
            is_ai_suggested=True,
            ai_confidence_score=Decimal('0.75'),
            ai_suggestion_reason=ai_reason,
            created_by=data_source.user 
        )
            if created:
                adjustments_suggested_count += 1
                suggested_adjustments_list.append(adjustment)

    data_source.processed_adjustments_count = adjustments_suggested_count
    data_source.save()

    return suggested_adjustments_list

def generate_ai_explanation(item_type, item_description):
    """Placeholder function to simulate AI explanation generation."""
    return f"AI suggests this is a {item_type} based on the description: \'{item_description}\'. This typically represents a non-recurring or owner-related expense for QoE purposes."

def generate_recast_pl(data_source_id, adjustments_ids, target_date):

    try:
        profit_loss_record = ProfitLossRecord.objects.filter(
            data_source_id=data_source_id,
            date__year=target_date.year,
            date__month=target_date.month
        ).first()

        if not profit_loss_record:
            raise ValueError("P&L record not found for the specified data source and period.")

        approved_adjustments = Adjustment.objects.filter(
            id__in=adjustments_ids, 
            status='APPROVED'
        )

        recast_revenue = profit_loss_record.revenue
        recast_cogs = profit_loss_record.cogs
        recast_gross_profit = profit_loss_record.gross_profit
        recast_operating_expenses = profit_loss_record.operating_expenses
        recast_ebitda = profit_loss_record.ebitda

        for adj in approved_adjustments:
            if adj.impact_type == 'POSITIVE':
                recast_ebitda += adj.amount
            elif adj.impact_type == 'NEGATIVE':
                recast_ebitda -= adj.amount
 
        recast_pl_data = {
            'date': target_date,
            'revenue': recast_revenue,
            'cogs': recast_cogs,
            'gross_profit': recast_gross_profit,
            'operating_expenses': recast_operating_expenses,
            'ebitda': recast_ebitda,
        }

        recast_pl_instance, created = RecastPL.objects.update_or_create(
            date=target_date, 
            defaults=recast_pl_data
        )
        recast_pl_instance.adjustments.set(approved_adjustments) 
        recast_pl_instance.save()

        print("Linked adjustments:", recast_pl_instance.adjustments.all())

        return recast_pl_instance

    except Exception as e:
        raise Exception(f"Error generating Recast P&L: {str(e)}")

def extract_pdf_data(file_path):

    print(f"Extracting data from PDF: {file_path}")

    return {"status": "success", "content": "Parsed content of PDF"}

def process_stoc_file(data_source):
    file_path = data_source.file_path.path
    uploaded_by_user = data_source.user
    processed_count = 0
    print(f"DEBUG: Entering process_stoc_file for data source {data_source.id}")  
    try:
        file_extension = os.path.splitext(file_path)[1].lower()
        df = None

        if file_extension == '.csv':
            df = pd.read_csv(file_path, header=None, sep=',', encoding='latin-1')
        elif file_extension in ['.xlsx', '.xls']:
            df = pd.read_excel(file_path, header=None)
        else:
            raise ValueError(f"Unsupported file type for STOC: {file_extension}. Only .csv and .xlsx/.xls are supported.")

        if df is None or df.empty:
            raise ValueError("Could not read STOC file or file is empty.")

        project_name = df.iloc[13, 1] if pd.notna(df.iloc[13, 1]) else None  
        location_city = df.iloc[14, 1] if pd.notna(df.iloc[14, 1]) else None 
        location_state = df.iloc[15, 1] if pd.notna(df.iloc[15, 1]) else None 
        
        external_financial_statement = df.iloc[18, 1] if pd.notna(df.iloc[18, 1]) else None
        calendar_year_fiscal_year = df.iloc[20, 1] if pd.notna(df.iloc[20, 1]) else None
        
        def parse_stoc_date(date_str):
            if pd.isna(date_str): return None
            date_str = str(date_str).strip()
            if not date_str: return None
            try:
                return pd.to_datetime(date_str).date()
            except ValueError:
                for fmt in ["%B %d, %Y", "%B %d,%Y", "%m/%d/%Y", "%Y-%m-%d"]:
                    try:
                        return datetime.strptime(date_str, fmt).date()
                    except ValueError:
                        pass
            return None

        stub_period = parse_stoc_date(df.iloc[21, 1]) 
        
        financial_reporting_period_2 = parse_stoc_date(df.iloc[22, 1])

        financial_reporting_period_1 = parse_stoc_date(df.iloc[23, 1])


        qe_summary_status = df.iloc[26, 1] if pd.notna(df.iloc[26, 1]) else 'N/A'
        recast_reported_status = df.iloc[27, 1] if pd.notna(df.iloc[27, 1]) else 'N/A' 
        recast_mgmt_adjusted_status = df.iloc[28, 1] if pd.notna(df.iloc[28, 1]) else 'N/A' 
        other_recast_status = df.iloc[29, 1] if pd.notna(df.iloc[29, 1]) else 'N/A'
        lead_profit_and_loss_status = df.iloc[30, 1] if pd.notna(df.iloc[30, 1]) else 'N/A' 
        monthly_profit_and_loss_status = df.iloc[31, 1] if pd.notna(df.iloc[31, 1]) else 'N/A'

 
        StocAccountingData.objects.create(
            project_name=project_name,
            location_city=location_city,
            location_state=location_state,
            external_financial_statement=external_financial_statement,
            calendar_year_fiscal_year=calendar_year_fiscal_year,
            stub_period=stub_period,
            financial_reporting_period_1=financial_reporting_period_1,
            financial_reporting_period_2=financial_reporting_period_2,
            qe_summary_status=qe_summary_status,
            recast_reported_status=recast_reported_status,
            recast_mgmt_adjusted_status=recast_mgmt_adjusted_status,
            other_recast_status=other_recast_status,
            lead_profit_and_loss_status=lead_profit_and_loss_status,
            monthly_profit_and_loss_status=monthly_profit_and_loss_status,
            uploaded_by=uploaded_by_user,
        )
        processed_count = 1 

        data_source.processed_records_count = processed_count
        data_source.save()

    except Exception as e:
        print(f"DEBUG: Error in process_stoc_file: {str(e)}")


def extract_date_from_payment_columns(row):
    record_date = None
    for i in range(1, 28): 
        date_col_regex = rf'payment\s* {re.escape(str(i))}\s* check date' 
        found_date_col = None

        for col_name in row.index:
            if re.search(date_col_regex, str(col_name).lower()):
                found_date_col = col_name
                break

        if found_date_col and pd.notna(row[found_date_col]):
            raw_date_val = row[found_date_col]
            try:
                record_date = pd.to_datetime(raw_date_val, errors='coerce')
                if pd.notna(record_date):
                    break

                for fmt in ["%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y", "%m-%d-%y"]:
                    try:
                        record_date = datetime.strptime(str(raw_date_val).strip(), fmt).date()
                        break
                    except ValueError:
                        pass
                if record_date:
                    break

            except Exception as e:
                pass
        else:
            pass

    if not record_date:
        pass
    return record_date 

def apply_header_formatting(ws, section):
    ws['A1'] = f"Recast P&L - {section.name}"
    ws['A1'].font = openpyxl.styles.Font(bold=True, size=16)
    ws.merge_cells('A1:D1')
    ws['A2'] = "Date: [Dynamically populated]"

def apply_financial_data_formatting(ws, section):
    ws['A4'] = f"Financial Data - {section.name}"
    ws['A4'].font = openpyxl.styles.Font(bold=True)
    ws['A5'] = "Revenue"
    ws['B5'] = "COGS"
    ws['C5'] = "Gross Profit"
    ws['D5'] = "Operating Expenses"
    ws['E5'] = "EBITDA"
    for col in ['A', 'B', 'C', 'D', 'E']:
        ws[f'{col}5'].fill = openpyxl.styles.PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

def apply_adjustments_formatting(ws, section):
    ws['A7'] = f"Adjustments - {section.name}"
    ws['A7'].font = openpyxl.styles.Font(bold=True)
    ws['A8'] = "Description"
    ws['B8'] = "Type"
    ws['C8'] = "Amount"
    ws['D8'] = "Rationale"
    ws['E8'] = "Source Reference"
    for col in ['A', 'B', 'C', 'D', 'E']:
        ws[f'{col}8'].fill = openpyxl.styles.PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

def fill_recast_pl_data(ws, recast_pl):
 
    ws['A2'] = f"Date: {recast_pl.date.strftime('%Y-%m')}"
    ws['A6'] = recast_pl.revenue
    ws['B6'] = recast_pl.cogs
    ws['C6'] = recast_pl.gross_profit
    ws['D6'] = recast_pl.operating_expenses
    ws['E6'] = recast_pl.ebitda

    current_row = 9
    for adjustment in recast_pl.adjustments.all():
        ws[f'A{current_row}'] = adjustment.description
        ws[f'B{current_row}'] = adjustment.adjustment_type
        ws[f'C{current_row}'] = adjustment.amount
        ws[f'D{current_row}'] = adjustment.rationale
        ws[f'E{current_row}'] = adjustment.source_reference
        current_row += 1

def export_recast_pl(recast_pl, template):
    wb = Workbook()
    ws = wb.active


    class MockSection:
        def __init__(self, name):
            self.name = name

    mock_template_sections = [
        MockSection("Header"),
        MockSection("Financial Data"),
        MockSection("Adjustments"),
    ]

    for section in mock_template_sections:
        if section.name == "Header":
            apply_header_formatting(ws, section)
        elif section.name == "Financial Data":
            apply_financial_data_formatting(ws, section)
        elif section.name == "Adjustments":
            apply_adjustments_formatting(ws, section)

    fill_recast_pl_data(ws, recast_pl)
    file_buffer = io.BytesIO()
    wb.save(file_buffer)
    file_buffer.seek(0)

    return file_buffer 